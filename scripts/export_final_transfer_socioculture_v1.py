#!/usr/bin/env python3
"""SOCIOCULTURE FINAL_TRANSFER V1: export / package / verify the Phase 2C KEEP set.

Selection source (only): data/phase2/SOCIOCULTURE/final_content_asset_index.jsonl
Question source: data/questions/SOCIOCULTURE/SOCIOCULTURE_questions.jsonl (+ images/)
Phase 2C metadata: data/phase2/SOCIOCULTURE/phase2c_final_curation.jsonl (read-only; used for
per-question curation fields and the DROP-contamination cross-check, never to re-select).

No re-classification, dedup, text/answer edits or image re-encoding. Images are byte copies.

  python3 scripts/export_final_transfer_socioculture_v1.py
"""
import collections
import datetime
import hashlib
import io
import json
import os
import re
import shutil
import sys
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SUBJECT = "SOCIOCULTURE"
PACKAGE = "SOCIOCULTURE_FINAL_TRANSFER_V1"
PKG_DIR_NAME = "SOCIOCULTURE_V1"
PHASE2C_COMMIT = "9e9c3e1"

SEL = "data/phase2/SOCIOCULTURE/final_content_asset_index.jsonl"
CUR = "data/phase2/SOCIOCULTURE/phase2c_final_curation.jsonl"
QSRC = "data/questions/SOCIOCULTURE/SOCIOCULTURE_questions.jsonl"
OUT_BASE = os.path.join(ROOT, "FINAL_TRANSFER")
OUT_DIR = os.path.join(OUT_BASE, PKG_DIR_NAME)
OUT_ZIP = os.path.join(OUT_BASE, PACKAGE + ".zip")
ZIP_DATE = (2026, 1, 1, 0, 0, 0)

XLSX_COLS = [
    "question_id", "exam_id", "year", "session", "question_number", "points", "official_answer",
    "primary_domain", "curriculum_codes", "required_content", "candidate_source",
    "adaptation_required", "content_relevance_reason", "adaptation_note",
    "phase1_review_required", "phase1_review_reasons", "phase2a_review_required",
    "phase2b_review_required", "phase2c_review_required", "phase2c_review_reasons",
    "stem", "passage", "choice_1", "choice_2", "choice_3", "choice_4", "choice_5",
    "image_file", "source_image_sha256",
]
WIDTHS = {"question_id": 30, "exam_id": 26, "stem": 50, "passage": 60, "content_relevance_reason": 50,
          "adaptation_note": 36, "required_content": 36, "curriculum_codes": 22, "primary_domain": 18,
          "image_file": 44, "source_image_sha256": 24, "phase1_review_reasons": 24,
          "phase2c_review_reasons": 24, "candidate_source": 18}


def sha_bytes(b):
    return hashlib.sha256(b).hexdigest()


def sha_file(p):
    with open(p, "rb") as f:
        return sha_bytes(f.read())


def jl(rel):
    with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
        return [json.loads(l) for l in f if l.strip()]


def join(xs):
    return " | ".join(str(x) for x in xs)


def build_records():
    sel = jl(SEL)
    q = {r["question_id"]: r for r in jl(QSRC)}
    cur = {r["question_id"]: r for r in jl(CUR)}
    recs = []
    for s in sel:
        qid = s["question_id"]
        p1, c = q[qid], cur[qid]
        src_img = os.path.join(ROOT, p1["question_image"])
        recs.append({
            "subject": p1["subject"],
            "question_id": qid,
            "exam_id": p1["exam_id"],
            "year": p1["year"],
            "session": p1["exam_type"],
            "question_number": p1["question_number"],
            "stem": p1["stem"],
            "passage": p1["passage"],
            "choices": p1["choices"],
            "official_answer": p1["official_answer"],
            "points": p1["points"],
            "image_file": "images/" + os.path.basename(p1["question_image"]),
            "primary_domain": s["primary_domain"],
            "curriculum_codes": s["curriculum_codes"],
            "required_content": c["required_content"],
            "candidate_source": s["candidate_source"],
            "adaptation_required": s["adaptation_required"],
            "content_relevance_reason": c["content_relevance_reason"],
            "adaptation_note": c["adaptation_note"],
            "phase1_review_required": p1["review_required"],
            "phase1_review_reasons": p1["review_reasons"],
            "phase2a_review_required": s["phase2a_review_required"],
            "phase2b_review_required": s["phase2b_review_required"],
            "phase2c_review_required": s["phase2c_review_required"],
            "phase2c_review_reasons": c["review_reasons"],
            "source_image_sha256": sha_file(src_img),
            "_src_img": src_img,
        })
    return sel, q, cur, recs


def write_xlsx(recs, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "questions"
    ws.append(XLSX_COLS)
    for r in recs:
        ch = {c["number"]: c["text"] for c in r["choices"]}
        row = []
        for col in XLSX_COLS:
            if col.startswith("choice_"):
                v = ch.get(int(col[-1]), "")
            else:
                v = r[col]
            row.append(join(v) if isinstance(v, list) else v)
        ws.append(row)
    for i, col in enumerate(XLSX_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 12)
        ws.cell(1, i).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.properties.created = datetime.datetime(2026, 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    # openpyxl stamps save time into core.xml and zip entries; normalize for determinism
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(path, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>2026-01-01T00:00:00Z\g<2>", data)
            zi = zipfile.ZipInfo(info.filename, ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)


README = """SOCIOCULTURE FINAL_TRANSFER V1

- Source pool: {pool:,} questions
- Phase2C candidates: {cand}
- Final KEEP: {keep}
- DROP: {drop}

Selection source:
final_content_asset_index.jsonl

Package:
- questions.jsonl
- questions.xlsx
- images/
- manifest.json

Important:
- All {keep} KEEP questions are included.
- No deduplication was performed.
- review_required items are intentionally retained.
- PNG files are byte-identical copies of Phase1 source images.
- Official answers come from Phase1 official answer linkage.
- No answer was inferred during export.
- This package is a source-content package for later Integrated Social Studies work.
- It is not itself a finished Integrated Social Studies question bank.
"""


def main():
    sel, q, cur, recs = build_records()
    if os.path.isdir(OUT_DIR):
        shutil.rmtree(OUT_DIR)
    os.makedirs(os.path.join(OUT_DIR, "images"))

    for r in recs:
        shutil.copyfile(r["_src_img"], os.path.join(OUT_DIR, r["image_file"]))
    with open(os.path.join(OUT_DIR, "questions.jsonl"), "w", encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps({k: v for k, v in r.items() if not k.startswith("_")},
                               ensure_ascii=False) + "\n")
    write_xlsx(recs, os.path.join(OUT_DIR, "questions.xlsx"))

    # manifest counts are computed from the written package files
    exported = [json.loads(l) for l in open(os.path.join(OUT_DIR, "questions.jsonl"), encoding="utf-8")]
    img_files = sorted(os.listdir(os.path.join(OUT_DIR, "images")))
    ncur = collections.Counter(r["final_decision"] for r in cur.values())
    manifest = {
        "package": PACKAGE,
        "subject": SUBJECT,
        "selection_source": SEL,
        "phase2c_commit": PHASE2C_COMMIT,
        "source_pool_questions": len(q),
        "phase2c_candidates": len(cur),
        "phase2c_keep": ncur["KEEP"],
        "phase2c_drop": ncur["DROP"],
        "questions": len(exported),
        "images": len(img_files),
        "official_answers_non_null": sum(r["official_answer"] is not None for r in exported),
        "adaptation": {k: sum(r["adaptation_required"] == k for r in exported)
                       for k in ("NONE", "LIGHT", "SUBSTANTIAL")},
        "phase2c_review_required_keep": sum(r["phase2c_review_required"] for r in exported),
        "phase1_review_required_keep": sum(r["phase1_review_required"] for r in exported),
        "by_candidate_source": dict(sorted(collections.Counter(r["candidate_source"] for r in exported).items())),
        "by_primary_domain": dict(collections.Counter(r["primary_domain"] for r in exported).most_common()),
        "hash_algorithm": "sha256",
        "files": {
            "questions.jsonl": sha_file(os.path.join(OUT_DIR, "questions.jsonl")),
            "images": {n: sha_file(os.path.join(OUT_DIR, "images", n)) for n in img_files},
        },
        "created_by": "SOCIOCULTURE FINAL_TRANSFER V1 exporter",
    }
    with open(os.path.join(OUT_DIR, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
        f.write("\n")
    with open(os.path.join(OUT_DIR, "README.txt"), "w", encoding="utf-8") as f:
        f.write(README.format(pool=len(q), cand=len(cur), keep=ncur["KEEP"], drop=ncur["DROP"]))

    # ZIP: single top-level dir, fixed timestamps, stored (PNG already compressed)
    if os.path.exists(OUT_ZIP):
        os.remove(OUT_ZIP)
    names = ["README.txt", "manifest.json", "questions.jsonl", "questions.xlsx"] + \
            ["images/" + n for n in img_files]
    with zipfile.ZipFile(OUT_ZIP, "w") as z:
        for n in names:
            zi = zipfile.ZipInfo(PKG_DIR_NAME + "/" + n, ZIP_DATE)
            zi.compress_type = zipfile.ZIP_STORED if n.endswith(".png") else zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            with open(os.path.join(OUT_DIR, n), "rb") as f:
                z.writestr(zi, f.read())

    qa = run_qa(sel, q, cur, recs)
    print(json.dumps(qa, ensure_ascii=False, indent=1))
    if not qa["PASS"]:
        sys.exit(1)


def run_qa(sel, q, cur, recs):
    qa = {}
    sel_ids = [s["question_id"] for s in sel]
    jrecs = [json.loads(l) for l in open(os.path.join(OUT_DIR, "questions.jsonl"), encoding="utf-8")]
    j_ids = [r["question_id"] for r in jrecs]
    ws = openpyxl.load_workbook(os.path.join(OUT_DIR, "questions.xlsx"), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    x_ids = [r[0] for r in rows[1:]]
    img_files = sorted(os.listdir(os.path.join(OUT_DIR, "images")))
    i_ids = [n[:-4] for n in img_files if n.endswith(".png")]
    manifest = json.load(open(os.path.join(OUT_DIR, "manifest.json"), encoding="utf-8"))
    S = set(sel_ids)

    qa["selection_records"] = len(sel_ids)
    qa["selection_unique_ids"] = len(S)
    qa["jsonl_records"] = len(jrecs)
    qa["excel_rows"] = len(x_ids)
    qa["excel_header_ok"] = list(rows[0]) == XLSX_COLS
    qa["png_files"] = len(img_files)
    qa["manifest_questions"] = manifest["questions"]
    qa["manifest_images"] = manifest["images"]
    ids = {"jsonl": j_ids, "excel": x_ids, "images": i_ids}
    qa["duplicate_ids"] = sum(len(v) - len(set(v)) for v in [sel_ids] + list(ids.values()))
    qa["missing_ids"] = sum(len(S - set(v)) for v in ids.values())
    qa["extra_ids"] = sum(len(set(v) - S) for v in ids.values())
    qa["drop_contamination"] = sum(cur[i]["final_decision"] != "KEEP" for i in j_ids)
    qa["phase2c_keep_missing_from_export"] = len({i for i, c in cur.items() if c["final_decision"] == "KEEP"} - set(j_ids))

    ans = [r["official_answer"] for r in jrecs]
    qa["answer_null"] = sum(a is None for a in ans)
    qa["answer_outside_1_5"] = sum(a is not None and a not in (1, 2, 3, 4, 5) for a in ans)
    qa["answer_mismatch_vs_phase1"] = sum(r["official_answer"] != q[r["question_id"]]["official_answer"] for r in jrecs)
    qa["text_mismatch_vs_phase1"] = sum(any(r[k] != q[r["question_id"]][k] for k in ("stem", "passage", "choices", "points"))
                                        for r in jrecs)
    qa["official_answers_non_null"] = len(ans) - qa["answer_null"]

    qa["metadata_errors"] = sum(not (
        r["subject"] == SUBJECT and r["question_id"] and r["exam_id"]
        and isinstance(r["question_number"], int) and r["question_number"] > 0
        and r["image_file"] and r["primary_domain"] and r["curriculum_codes"]
        and r["content_relevance_reason"] and r["adaptation_required"] in ("NONE", "LIGHT")) for r in jrecs)
    qa["substantial_keep"] = sum(r["adaptation_required"] == "SUBSTANTIAL" for r in jrecs)

    mism = 0
    for r in recs:
        exp = sha_file(os.path.join(OUT_DIR, r["image_file"]))
        mism += not (sha_file(r["_src_img"]) == exp == r["source_image_sha256"]
                     == q[r["question_id"]]["question_image_sha256"])
    qa["source_to_export_checked"] = len(recs)
    qa["source_to_export_mismatch"] = mism

    with zipfile.ZipFile(OUT_ZIP) as z:
        qa["zip_corrupt_files"] = 0 if z.testzip() is None else 1
        zn = z.namelist()
        tops = {n.split("/")[0] for n in zn}
        qa["zip_top_level_dirs"] = sorted(tops)
        zj = [json.loads(l) for l in z.read(PKG_DIR_NAME + "/questions.jsonl").decode("utf-8").splitlines() if l]
        qa["zip_jsonl_records"] = len(zj)
        zimgs = [n for n in zn if n.startswith(PKG_DIR_NAME + "/images/") and n.endswith(".png")]
        qa["zip_images"] = len(zimgs)
        zm = 0
        for n in img_files:
            zm += sha_bytes(z.read(PKG_DIR_NAME + "/images/" + n)) != sha_file(os.path.join(OUT_DIR, "images", n))
        qa["export_to_zip_checked"] = len(img_files)
        qa["export_to_zip_mismatch"] = zm
        qa["zip_xlsx_rows"] = openpyxl.load_workbook(io.BytesIO(z.read(PKG_DIR_NAME + "/questions.xlsx")),
                                                     read_only=True).active.max_row - 1
        qa["zip_manifest_questions"] = json.loads(z.read(PKG_DIR_NAME + "/manifest.json"))["questions"]
        qa["zip_other_files_identical"] = all(
            z.read(PKG_DIR_NAME + "/" + n) == open(os.path.join(OUT_DIR, n), "rb").read()
            for n in ("README.txt", "manifest.json", "questions.jsonl", "questions.xlsx"))
    qa["zip_size_bytes"] = os.path.getsize(OUT_ZIP)

    n = qa["selection_records"]
    qa["PASS"] = all([
        n == qa["selection_unique_ids"] == qa["jsonl_records"] == qa["excel_rows"] == qa["png_files"]
        == qa["manifest_questions"] == qa["manifest_images"] == qa["official_answers_non_null"]
        == qa["zip_jsonl_records"] == qa["zip_images"] == qa["zip_xlsx_rows"] == qa["zip_manifest_questions"],
        qa["excel_header_ok"], qa["zip_other_files_identical"],
        qa["zip_top_level_dirs"] == [PKG_DIR_NAME],
        not any(qa[k] for k in ("duplicate_ids", "missing_ids", "extra_ids", "drop_contamination",
                                "phase2c_keep_missing_from_export", "answer_null", "answer_outside_1_5",
                                "answer_mismatch_vs_phase1", "text_mismatch_vs_phase1", "metadata_errors",
                                "substantial_keep", "source_to_export_mismatch", "export_to_zip_mismatch",
                                "zip_corrupt_files")),
    ])
    return qa


if __name__ == "__main__":
    main()
